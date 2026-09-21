from typing import List, Dict, Optional, Any
from docx import Document
from docx.shared import Cm
from docx.enum.text import WD_COLOR_INDEX
import re


# ==========================================
# Универсальный хелпер для извлечения полей
# ==========================================
def _get_field(obj: Any, key: str, default: Any = None) -> Any:
    """
    Универсальное извлечение поля из объекта.
    Работает с датаклассами (CheckError), обычными объектами и словарями.
    """
    if isinstance(obj, dict):
        return obj.get(key, default)
    return getattr(obj, key, default)


class DocumentMarker:
    """
    Модуль для визуального выделения ошибок в документе.
    
    Применение:
        marker = DocumentMarker("input.docx", errors)
        stats = marker.mark_and_save("output.docx")
    
    Результат:
        - Проблемные места выделены цветом
        - В конец документа добавлена таблица-отчёт
    """
    
    # Цветовая схема по умолчанию
    DEFAULT_COLORS = {
        'error': WD_COLOR_INDEX.PINK,
        'warning': WD_COLOR_INDEX.YELLOW,
        'info': WD_COLOR_INDEX.TURQUOISE,
    }
    
    def __init__(self, input_path: str, errors: List[Any], colors: Dict = None):
        """
        Args:
            input_path: Путь к исходному .docx файлу
            errors: Список ошибок от чекеров (объекты CheckError или словари)
            colors: Опциональная кастомная цветовая схема
        """
        self.input_path = input_path
        self.doc = Document(input_path)
        self.errors = errors
        self.colors = colors or self.DEFAULT_COLORS
        
        # Статистика обработки
        self.stats = {
            'total': len(errors),
            'marked': 0,
            'not_found': 0,
            'skipped': 0,
        }
    
    def mark_and_save(self, output_path: str, add_report: bool = True) -> Dict:
        """
        Применяет маркеры и сохраняет документ.
        """
        for error in self.errors:
            self._mark_error(error)
        
        if add_report and self.errors:
            self._add_report_section()
        
        self.doc.save(output_path)
        return self.stats
    
    # ==========================================
    # Вспомогательные методы
    # ==========================================
    
    def _get_severity_str(self, error: Any) -> str:
        """Преобразует severity в строку (обрабатывает Enum и строки)"""
        severity = _get_field(error, 'severity', 'info')
        # Обработка Enum (у Severity.value есть строка)
        if hasattr(severity, 'value'):
            return str(severity.value).lower()
        return str(severity).lower()
    
    def _get_color(self, error: Any) -> WD_COLOR_INDEX:
        """Получает цвет выделения для ошибки"""
        severity_str = self._get_severity_str(error)
        return self.colors.get(severity_str, WD_COLOR_INDEX.YELLOW)
    
    def _extract_text_from_context(self, context: str) -> Optional[str]:
        """
        Извлекает текст элемента из контекста ошибки.
        
        Примеры:
            "Заголовок: 'Требования к насосам'" → "Требования к насосам"
            "Абзац: 'Текст абзаца...'" → "Текст абзаца"
            "после: 'Текст1...', перед: 'Текст2...'" → "Текст1"
        """
        matches = re.findall(r"'([^']+)'", context)
        if not matches:
            return None
        
        text = max(matches, key=len)
        text = re.sub(r'\.{2,}$', '', text).strip()
        
        if len(text) < 3:
            return None
        
        return text
    
    def _extract_table_index(self, context: str) -> Optional[int]:
        """
        Извлекает номер таблицы из контекста (0-based index).
        """
        match = re.search(r'Таблица\s*#(\d+)', context)
        if match:
            return int(match.group(1)) - 1
        return None
    
    def _find_paragraphs_by_text(self, text: str) -> List:
        """
        Ищет параграфы, содержащие указанный текст.
        """
        results = []
        search_text = text.strip()
        
        for para in self.doc.paragraphs:
            para_text = para.text.strip()
            if search_text in para_text or para_text in search_text:
                results.append(para)
        
        return results
    
    # ==========================================
    # Методы выделения
    # ==========================================
    
    def _highlight_paragraph(self, paragraph, color: WD_COLOR_INDEX):
        """Выделяет параграф цветом (все runs)"""
        for run in paragraph.runs:
            run.font.highlight_color = color
    
    def _highlight_table(self, table, color: WD_COLOR_INDEX):
        """Выделяет всю таблицу цветом"""
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    self._highlight_paragraph(para, color)
    
    # ==========================================
    # Основной метод обработки ошибки
    # ==========================================
    
    def _mark_error(self, error: Any):
        """Выделяет одну ошибку цветом в документе"""
        node_type = _get_field(error, 'node_type', '')
        context = _get_field(error, 'context', '')
        color = self._get_color(error)
        
        # === НОВОЕ: Обработка node_type="llm_check" ===
        # Для LLM-чекеров пытаемся найти параграф по контексту
        if node_type == 'llm_check':
            text = self._extract_text_from_context(context)
            if text:
                paragraphs = self._find_paragraphs_by_text(text)
                if paragraphs:
                    for para in paragraphs:
                        self._highlight_paragraph(para, color)
                    self.stats['marked'] += 1
                else:
                    # Если не нашли, пропускаем (не выделяем весь документ)
                    self.stats['not_found'] += 1
            else:
                self.stats['skipped'] += 1
            return
        
        # === Остальная логика без изменений ===
        # --- Параграфы, заголовки, пункты списков ---
        if node_type in ('heading', 'paragraph', 'list_item'):
            text = self._extract_text_from_context(context)
            if text:
                paragraphs = self._find_paragraphs_by_text(text)
                if paragraphs:
                    for para in paragraphs:
                        self._highlight_paragraph(para, color)
                    self.stats['marked'] += 1
                else:
                    self.stats['not_found'] += 1
            else:
                self.stats['skipped'] += 1
        
        # --- Таблицы ---
        elif node_type == 'table':
            table_idx = self._extract_table_index(context)
            if table_idx is not None:
                tables = self.doc.tables
                if 0 <= table_idx < len(tables):
                    self._highlight_table(tables[table_idx], color)
                    self.stats['marked'] += 1
                else:
                    self.stats['not_found'] += 1
            else:
                self.stats['skipped'] += 1
        
        # --- Документные ошибки ---
        elif node_type == 'document':
            self.stats['skipped'] += 1
        
        # --- Сноски ---
        elif node_type == 'footnote':
            self.stats['skipped'] += 1
        
        # --- Неизвестный тип ---
        else:
            self.stats['skipped'] += 1
        
    # ==========================================
    # Генерация отчёта
    # ==========================================
    
    def _add_report_section(self):
        """Добавляет раздел с отчётом об ошибках в конец документа"""
        self.doc.add_page_break()
        
        # --- Заголовок ---
        self.doc.add_heading('Отчёт нормоконтроля', level=1)
        
        # --- Легенда цветов ---
        legend = self.doc.add_paragraph()
        legend.add_run('Легенда: ').bold = True
        run_err = legend.add_run('■ ERROR  ')
        run_err.font.highlight_color = WD_COLOR_INDEX.PINK
        run_warn = legend.add_run('■ WARNING  ')
        run_warn.font.highlight_color = WD_COLOR_INDEX.YELLOW
        run_info = legend.add_run('■ INFO')
        run_info.font.highlight_color = WD_COLOR_INDEX.TURQUOISE
        
        # --- Статистика ---
        errors_count = sum(1 for e in self.errors if self._get_severity_str(e) == 'error')
        warnings_count = sum(1 for e in self.errors if self._get_severity_str(e) == 'warning')
        info_count = sum(1 for e in self.errors if self._get_severity_str(e) == 'info')
        
        stats_para = self.doc.add_paragraph()
        stats_para.add_run(f'Всего обнаружено: {self.stats["total"]} замечаний\n').bold = True
        stats_para.add_run(f'  • Ошибки (ERROR): {errors_count}\n')
        stats_para.add_run(f'  • Предупреждения (WARNING): {warnings_count}\n')
        stats_para.add_run(f'  • Информация (INFO): {info_count}\n')
        stats_para.add_run(f'\nВыделено в документе: {self.stats["marked"]}\n')
        stats_para.add_run(f'Не найдено в документе: {self.stats["not_found"]}\n')
        
        self.doc.add_paragraph()
        
        # --- Таблица с ошибками ---
        table = self.doc.add_table(rows=1, cols=6)
        table.style = 'Table Grid'
        
        # Заголовки таблицы
        header_cells = table.rows[0].cells
        headers = ['№', 'Тип', 'Серьёзность', 'Описание', 'ГОСТ', 'Контекст']
        for i, header in enumerate(headers):
            header_cells[i].text = header
            for run in header_cells[i].paragraphs[0].runs:
                run.font.bold = True
        
        # Строки с ошибками
        for idx, error in enumerate(self.errors, 1):
            row_cells = table.add_row().cells
            row_cells[0].text = str(idx)
            row_cells[1].text = _get_field(error, 'node_type', '')
            row_cells[2].text = self._get_severity_str(error).upper()
            row_cells[3].text = _get_field(error, 'message', '')
            row_cells[4].text = _get_field(error, 'gost_ref', '')
            row_cells[5].text = _get_field(error, 'context', '')
            
            # Цветовая маркировка ячейки "Серьёзность"
            severity_str = self._get_severity_str(error)
            color = self.colors.get(severity_str, WD_COLOR_INDEX.YELLOW)
            for para in row_cells[2].paragraphs:
                for run in para.runs:
                    run.font.highlight_color = color
        
        # Устанавливаем ширину колонок
        for row in table.rows:
            row.cells[0].width = Cm(1)
            row.cells[1].width = Cm(2)
            row.cells[2].width = Cm(2.5)
            row.cells[3].width = Cm(6)
            row.cells[4].width = Cm(3)
            row.cells[5].width = Cm(4)


# ==========================================
# Функция-обёртка для удобства
# ==========================================

def mark_document(input_path: str, output_path: str, errors: List[Any], 
                  add_report: bool = True) -> Dict:
    """
    Быстрая функция для выделения ошибок в документе.
    
    Args:
        input_path: Путь к исходному .docx
        output_path: Путь для сохранения
        errors: Список ошибок от чекеров (объекты CheckError)
        add_report: Добавить отчёт в конец документа
        
    Returns:
        Статистика обработки
    """
    marker = DocumentMarker(input_path, errors)
    return marker.mark_and_save(output_path, add_report=add_report)